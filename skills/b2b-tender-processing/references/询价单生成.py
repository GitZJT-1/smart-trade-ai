# -*- coding: utf-8 -*-
"""
NLMK 招标 → 标准化询价单 Excel 生成器（可复用）。
与 SKILL.md Step 4 对应，格式已与模板（260818-20155582齿轮支座-待报价）逐列核对。

用法：改 tenders 里的行项数据后运行。
  - folder  = 文件夹名 `{日期}-{标段}{中文品名}-{状态}`（状态：待报价/已报价）
  - fname   = 询价单文件名 `{标段}{中文品名}等询价单.xlsx`
  - item 元组 = (物料号[str], 俄文名称, 中文名称, 数量[int], 材质[str])
    数量列填整数（俄文逗号=小数点需先换算，"1,000 ШТ" -> 1）；
    单重/总重 扫描件拿不到写「按图纸」；材质 从主 PDF 名称或图纸标题栏取，拿不到写「按图纸」；
    单价/总价/交货期 留空（公司内部上报，铁律）。
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

# 输出根目录：环境变量 TENDER_OUT_DIR 优先，否则当前用户的桌面
DESKTOP = os.environ.get("TENDER_OUT_DIR", os.path.join(os.path.expanduser("~"), "Desktop"))

# ===== 唯一要改的地方 =====
tenders = {
    "20150000": {   # 示例：换成实际标段号
        "folder": "260818-20150000分配头-待报价",
        "fname": "20150000分配头等询价单.xlsx",
        "items": [
            ("611170", "Шайба распределительная правая 100ДШ-01", "右分配盘 100ДШ-01", 1, "按图纸"),
            ("985535", "Шайба ячейковая Л63С.00.011 СЧ20 ГОСТ 1412-85", "格室盘 Л63С.00.011", 1, "СЧ20 ГОСТ 1412-85"),
        ],
    },
}

# 11 列表头（\n 在单元格内换行）
header = ["序号", "物料号", "名称", "中文名称", "数量\npcs", "图纸单重\nkg",
          "总重", "材质", "单价", "总价", "交货期\nLead time"]

thin = Side(style="thin")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_font = Font(name="Arial", size=11, bold=True)          # 表头加粗
data_font = Font(name="Times New Roman", size=11)
bold_font = Font(name="Times New Roman", size=11, bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

col_widths = {"A": 9.0, "B": 14.25, "C": 34.5, "D": 27.6, "E": 9.0,
              "F": 11.75, "G": 11.75, "H": 17.5, "I": 11.63, "J": 11.63, "K": 11.75}


def build(folder, fname, items):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # 表头（第 1 行）
    for col, h in enumerate(header, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font, c.alignment, c.border = hdr_font, center, border
    ws.row_dimensions[1].height = 30

    # 数据行（第 2 行起）
    first_data = 2
    for i, (mat, name_ru, name_cn, qty, material) in enumerate(items):
        r = first_data + i
        vals = [str(i + 1), mat, name_ru, name_cn, qty, "按图纸", "按图纸", material, None, None, None]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.font, c.alignment, c.border = data_font, center, border
        ws.row_dimensions[r].height = 32

    # 汇总行：A:C 合并写 total，E 列 =SUM(数量)，其余留空
    last = first_data + len(items) - 1
    tr = last + 1
    ws.cell(row=tr, column=1, value="total").font = bold_font
    ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=3)
    ws.cell(row=tr, column=1).alignment = center
    ws.cell(row=tr, column=5, value=f"=SUM(E{first_data}:E{last})")
    for col in range(1, 12):
        c = ws.cell(row=tr, column=col)
        c.border = border
        if col not in (1, 2, 3):
            c.font, c.alignment = data_font, center
    ws.row_dimensions[tr].height = 32

    for col_letter, w in col_widths.items():
        ws.column_dimensions[col_letter].width = w

    out = os.path.join(DESKTOP, folder, fname)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"SAVED: {out}  ({len(items)} items)")


if __name__ == "__main__":
    for folder, info in tenders.items():
        build(info["folder"], info["fname"], info["items"])
